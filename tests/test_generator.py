from load_test import make_xlsx
from openpyxl import load_workbook
import io


def test_generated_workbook_has_153_rows_and_10_columns():
    content = make_xlsx(1)
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    assert ws.max_row == 154  # header + 153 fund-allocation rows
    assert ws.max_column == 10
